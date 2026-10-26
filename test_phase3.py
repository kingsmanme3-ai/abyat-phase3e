import io

from openpyxl import load_workbook


def login(client, user="admin", pw="admin-pass-123"):
    r = client.post("/api/auth/login", data={"username": user, "password": pw})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def seed(client, h):
    site = client.post("/api/sites", json={"name": "موقع التقارير"}, headers=h).json()
    b = client.post("/api/buildings", json={"site_id": site["id"], "name": "ب1", "floors": 3}, headers=h).json()
    client.post("/api/units/bulk", json={"building_id": b["id"], "numbers": ["1", "2", "10", "3"]}, headers=h)
    units = {u["number"]: u for u in client.get(f"/api/units?site_id={site['id']}", headers=h).json()}
    client.patch(f"/api/units/{units['1']['id']}/status", headers=h,
                 json={"status": "rented", "tenant_name": "محمد", "contract_end": "2000-01-01"})
    client.patch(f"/api/units/{units['2']['id']}/status", headers=h,
                 json={"status": "rented", "tenant_name": "علي", "contract_end": "2999-01-01"})
    client.patch(f"/api/units/{units['3']['id']}/status", headers=h, json={"status": "maintenance"})
    return site, b


def test_report_types_and_roles(client):
    h = login(client)
    types = client.get("/api/reports/types", headers=h).json()
    keys = {t["key"]: t for t in types}
    assert keys["users"]["available"] and keys["payments"]["available"] is False

    client.post("/api/users", headers=h, json={"username": "viewer1", "full_name": "م", "password": "viewer-pass-1"})
    hv = login(client, "viewer1", "viewer-pass-1")
    vt = {t["key"] for t in client.get("/api/reports/types", headers=hv).json()}
    assert "users" not in vt and "units" in vt
    assert client.get("/api/reports/users", headers=hv).status_code == 403
    assert client.get("/api/reports/payments", headers=h).status_code == 404
    assert client.get("/api/report-schedules", headers=hv).status_code == 403


def test_reports_content(client):
    h = login(client)
    site, _ = seed(client, h)
    sid = site["id"]

    d = client.get(f"/api/reports/units?site_id={sid}", headers=h).json()
    assert d["total_rows"] == 4
    assert [r["unit"] for r in d["rows"]] == ["1", "2", "3", "10"]          # ترتيب طبيعي
    assert d["rows"][0]["status"] == "مؤجرة" and d["rows"][0]["tenant"] == "محمد"
    only = client.get(f"/api/reports/units?site_id={sid}&status=maintenance", headers=h).json()
    assert only["total_rows"] == 1

    p = client.get(f"/api/reports/properties?site_id={sid}", headers=h).json()
    row = p["rows"][0]
    assert row["total"] == 4 and row["rented"] == 2 and row["maintenance"] == 1 and row["occupancy"] == 50.0

    c = client.get(f"/api/reports/contracts?site_id={sid}", headers=h).json()
    assert c["total_rows"] == 2 and c["rows"][0]["state"] == "منتهٍ"
    soon = client.get(f"/api/reports/contracts?site_id={sid}&days=30", headers=h).json()
    assert soon["total_rows"] == 1 and soon["rows"][0]["tenant"] == "محمد"

    u = client.get("/api/reports/users", headers=h).json()
    assert any(r["username"] == "admin" and r["role"] == "مدير" for r in u["rows"])


def test_exports(client):
    h = login(client)
    r = client.get("/api/reports/units?format=xlsx", headers=h)
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert "filename*=UTF-8''" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.cell(1, 1).value == "الموقع" and ws.max_row >= 2

    r = client.get("/api/reports/contracts?format=csv", headers=h)
    assert r.status_code == 200 and r.content.startswith(b"\xef\xbb\xbf")


def test_schedules_lifecycle(client):
    h = login(client)
    # تحقق: أسبوعي بلا يوم
    bad = client.post("/api/report-schedules", headers=h, json={
        "name": "س", "report_type": "units", "frequency": "weekly"})
    assert bad.status_code == 422
    bad = client.post("/api/report-schedules", headers=h, json={
        "name": "س", "report_type": "units", "frequency": "daily", "recipients": ["not-an-email"]})
    assert bad.status_code == 422

    r = client.post("/api/report-schedules", headers=h, json={
        "name": "العقود الأسبوعي", "report_type": "contracts", "days": 60, "frequency": "weekly",
        "weekday": 0, "hour": 7, "fmt": "xlsx", "recipients": ["A@Example.com", "a@example.com"]})
    assert r.status_code == 201, r.text
    s = r.json()
    assert s["recipients"] == ["a@example.com"] and s["next_run_at"] and s["is_active"]

    # تشغيل فوري ثم تنزيل الملف
    run = client.post(f"/api/report-schedules/{s['id']}/run", headers=h).json()
    assert run["status"] == "ok" and run["trigger"] == "manual"
    assert "SMTP" in (run["email_note"] or "")            # لا إعدادات بريد في الاختبار
    dl = client.get(f"/api/report-runs/{run['id']}/download", headers=h)
    assert dl.status_code == 200 and load_workbook(io.BytesIO(dl.content)).active.cell(1, 1).value == "الموقع"
    assert client.get("/api/report-runs", headers=h).json()[0]["id"] == run["id"]

    # تعديل: يومي يلغي يوم الأسبوع، وإيقافه يمسح الموعد التالي وتفعيله يعيده
    u = client.patch(f"/api/report-schedules/{s['id']}", headers=h, json={"frequency": "daily"}).json()
    assert u["weekday"] is None and u["frequency"] == "daily"
    off = client.patch(f"/api/report-schedules/{s['id']}", headers=h, json={"is_active": False}).json()
    assert off["next_run_at"] is None
    on = client.patch(f"/api/report-schedules/{s['id']}", headers=h, json={"is_active": True}).json()
    assert on["next_run_at"]
    assert client.patch(f"/api/report-schedules/{s['id']}", headers=h,
                        json={"frequency": "monthly"}).status_code == 422   # يلزم يوم الشهر

    # الحذف يُبقي الملف المولّد متاحاً
    assert client.delete(f"/api/report-schedules/{s['id']}", headers=h).status_code == 204
    assert client.get(f"/api/report-runs/{run['id']}/download", headers=h).status_code == 200
    assert client.get("/api/report-schedules", headers=h).json() == []


def test_due_schedule_runs_once(client):
    from datetime import timedelta
    from sqlalchemy import select
    from app.database import SessionLocal
    from app.models import ReportRun, ReportSchedule, utcnow
    from app.report_runner import run_due_schedules

    h = login(client)
    s = client.post("/api/report-schedules", headers=h, json={
        "name": "مستحقة", "report_type": "properties", "frequency": "daily", "hour": 8}).json()
    with SessionLocal() as db:
        db.get(ReportSchedule, s["id"]).next_run_at = utcnow() - timedelta(hours=1)
        db.commit()
    assert run_due_schedules() == 1
    assert run_due_schedules() == 0                       # الموعد التالي صار في المستقبل
    with SessionLocal() as db:
        sch = db.get(ReportSchedule, s["id"])
        assert sch.last_run_at is not None
        runs = db.scalars(select(ReportRun).where(ReportRun.schedule_id == s["id"])).all()
        assert len(runs) == 1 and runs[0].trigger == "schedule" and runs[0].status == "ok"
