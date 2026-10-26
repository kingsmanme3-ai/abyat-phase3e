"""اختبارات نقية للمرحلة 3 (الجدولة وتوليد الملفات) بلا قاعدة بيانات."""
import io
from datetime import date, datetime, timezone

from openpyxl import load_workbook

from app.report_engine import to_csv, to_xlsx
from app.scheduling import compute_next_run

# الخميس 22-10-2026 الساعة 09:00 بتوقيت السعودية = 06:00 UTC
NOW = datetime(2026, 10, 22, 6, 0, tzinfo=timezone.utc)


def utc(*a):
    return datetime(*a, tzinfo=timezone.utc)


def test_daily():
    assert compute_next_run("daily", 10, None, None, NOW, 3) == utc(2026, 10, 22, 7)    # اليوم 10:00
    assert compute_next_run("daily", 8, None, None, NOW, 3) == utc(2026, 10, 23, 5)     # فات 8:00 فغداً


def test_weekly_sunday_based():
    assert compute_next_run("weekly", 8, 0, None, NOW, 3) == utc(2026, 10, 25, 5)       # الأحد القادم
    assert compute_next_run("weekly", 8, 4, None, NOW, 3) == utc(2026, 10, 29, 5)       # الخميس: فات موعده
    assert compute_next_run("weekly", 10, 4, None, NOW, 3) == utc(2026, 10, 22, 7)      # الخميس اليوم 10:00
    assert compute_next_run("weekly", 8, 6, None, NOW, 3) == utc(2026, 10, 24, 5)       # السبت


def test_monthly():
    assert compute_next_run("monthly", 8, None, 1, NOW, 3) == utc(2026, 11, 1, 5)
    assert compute_next_run("monthly", 8, None, 22, NOW, 3) == utc(2026, 11, 22, 5)
    assert compute_next_run("monthly", 8, None, 28, NOW, 3) == utc(2026, 10, 28, 5)
    assert compute_next_run("monthly", 8, None, 5, utc(2026, 12, 30, 6), 3) == utc(2027, 1, 5, 5)


def test_midnight_local_crosses_utc_day():
    # 00:00 بتوقيت +3 هو 21:00 UTC لليوم السابق
    assert compute_next_run("daily", 0, None, None, NOW, 3) == utc(2026, 10, 22, 21)


COLS = [("a", "الاسم"), ("d", "التاريخ"), ("n", "رقم")]
ROWS = [{"a": "محمد", "d": date(2027, 3, 1), "n": 5},
        {"a": "=HYPERLINK(\"http://x\")", "d": None, "n": None},
        {"a": "عادي\x07", "d": datetime(2026, 10, 22, 9, 30, tzinfo=timezone.utc), "n": 0}]


def test_csv_bom_and_formula_guard():
    raw = to_csv(COLS, ROWS)
    assert raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    assert text.splitlines()[0] == "الاسم,التاريخ,رقم"
    assert "2027-03-01" in text and "'=HYPERLINK" in text and "\x07" not in text


def test_xlsx_roundtrip_rtl_and_no_formula():
    wb = load_workbook(io.BytesIO(to_xlsx("تقرير/الوحدات:اختبار", COLS, ROWS)))
    ws = wb.active
    assert ws.sheet_view.rightToLeft is True and ws.freeze_panes == "A2"
    assert "/" not in ws.title and len(ws.title) <= 31
    assert [c.value for c in ws[1]] == ["الاسم", "التاريخ", "رقم"]
    assert ws["A2"].value == "محمد" and ws["B2"].value.date() == date(2027, 3, 1) and ws["C2"].value == 5
    assert ws["A3"].data_type == "s" and ws["A3"].value.startswith("=")
    assert ws["A4"].value == "عادي" and ws["B4"].value == datetime(2026, 10, 22, 9, 30)


def test_xlsx_empty_rows():
    wb = load_workbook(io.BytesIO(to_xlsx("فارغ", COLS, [])))
    assert wb.active.max_row == 1
