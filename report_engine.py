"""توليد ملفات التقارير (Excel/CSV) من أعمدة وصفوف. دوال نقية بلا اعتماد على قاعدة البيانات."""
import csv
import io
import re
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

Columns = list[tuple[str, str]]          # [(key, label)]

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_BAD_SHEET = re.compile(r"[\[\]:*?/\\]")
_FORMULA_START = ("=", "+", "-", "@")


def _is_number(s: str) -> bool:
    try:
        float(s)
        return True
    except ValueError:
        return False


def _csv_cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.replace(tzinfo=None).strftime("%Y-%m-%d %H:%M")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str):
        v = _ILLEGAL.sub("", v)
        # حماية من حقن الصيغ عند فتح الملف في Excel
        if v and v[0] in _FORMULA_START and not _is_number(v):
            v = "'" + v
    return v


def to_csv(columns: Columns, rows: list[dict]) -> bytes:
    """CSV بترميز UTF-8 مع BOM ليعرض Excel العربية صحيحة."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([label for _, label in columns])
    for r in rows:
        w.writerow([_csv_cell(r.get(k)) for k, _ in columns])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def to_xlsx(title: str, columns: Columns, rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = (_BAD_SHEET.sub(" ", title).strip() or "تقرير")[:31]
    ws.sheet_view.rightToLeft = True

    head_fill = PatternFill("solid", fgColor="0F766E")
    for i, (_, label) in enumerate(columns, start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    widths = [len(label) + 4 for _, label in columns]
    for ri, r in enumerate(rows, start=2):
        for ci, (k, _) in enumerate(columns, start=1):
            v = r.get(k)
            if isinstance(v, str):
                v = _ILLEGAL.sub("", v)
            elif isinstance(v, datetime):
                v = v.replace(tzinfo=None)
            cell = ws.cell(row=ri, column=ci, value=v)
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"           # لا يُفسَّر كصيغة
            elif isinstance(v, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(v, date):
                cell.number_format = "yyyy-mm-dd"
            if v is not None:
                widths[ci - 1] = max(widths[ci - 1], min(len(str(v)) + 2, 45))

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(w, 45))
    ws.freeze_panes = "A2"
    if columns:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
